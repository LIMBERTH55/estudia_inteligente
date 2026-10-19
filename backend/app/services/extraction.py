"""Extracción de texto de documentos: PDF, DOCX, PPTX, TXT e imágenes (OCR con IA)."""
from __future__ import annotations

import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

import pymupdf as fitz  # PyMuPDF

from ..ai import AIError, ImageInput, get_provider
from ..ai.prompts import OCR_SYSTEM, ocr_prompt

logger = logging.getLogger(__name__)

ProgressCb = Callable[[int, str], None]  # (porcentaje 0-100 del paso, descripción)

SUPPORTED_EXTENSIONS = {"pdf", "docx", "pptx", "txt", "jpg", "jpeg", "png"}
OCR_MIN_CHARS = 40  # una página con menos texto que esto se considera escaneada


@dataclass
class PageContent:
    page: int | None
    text: str
    section: str | None = None


@dataclass
class ExtractionResult:
    pages: list[PageContent] = field(default_factory=list)
    page_count: int = 0
    word_count: int = 0
    ocr_used: bool = False
    warnings: list[str] = field(default_factory=list)


def _count_words(text: str) -> int:
    return len(text.split())


def _ocr_image_bytes(data: bytes, mime_type: str, hint: str = "") -> str:
    """OCR mediante el modelo de visión del proveedor de IA configurado."""
    provider = get_provider()
    text = provider.generate_text(
        ocr_prompt(hint),
        system=OCR_SYSTEM,
        images=[ImageInput(data=data, mime_type=mime_type)],
        temperature=0.1,
    )
    if "[SIN_TEXTO_LEGIBLE]" in text:
        return ""
    return text.strip()


def _extract_pdf(path: Path, cb: ProgressCb | None) -> ExtractionResult:
    result = ExtractionResult()
    doc = fitz.open(path)
    total = doc.page_count
    result.page_count = total
    for i, page in enumerate(doc):
        if cb:
            cb(int((i / max(total, 1)) * 90), f"Extrayendo texto: página {i + 1} de {total}")
        text = page.get_text("text").strip()
        if len(text) < OCR_MIN_CHARS:
            # Página escaneada o con imagen: intentar OCR con IA
            try:
                pix = page.get_pixmap(dpi=150)
                ocr_text = _ocr_image_bytes(pix.tobytes("png"), "image/png", f"página {i + 1} de un PDF")
                if ocr_text:
                    text = ocr_text
                    result.ocr_used = True
                else:
                    result.warnings.append(f"Página {i + 1}: sin texto legible tras OCR.")
                    text = ""
            except AIError as exc:
                result.warnings.append(
                    f"Página {i + 1} parece escaneada y no se pudo aplicar OCR: {exc}"
                )
                text = ""
            except Exception as exc:  # noqa: BLE001
                logger.warning("OCR falló en página %s: %s", i + 1, exc)
                result.warnings.append(f"Página {i + 1}: falló el OCR.")
                text = ""
        if text:
            result.pages.append(PageContent(page=i + 1, text=text))
    doc.close()
    result.word_count = sum(_count_words(p.text) for p in result.pages)
    return result


def _extract_docx(path: Path, cb: ProgressCb | None) -> ExtractionResult:
    from docx import Document as DocxDocument
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    result = ExtractionResult()
    doc = DocxDocument(str(path))
    if cb:
        cb(20, "Leyendo documento de Word")
    section: str | None = None
    blocks: list[tuple[str | None, str]] = []  # (sección, texto)

    def append(text: str) -> None:
        if text.strip():
            blocks.append((section, text.strip()))

    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            para = Paragraph(child, doc)
            style = (para.style.name or "") if para.style is not None else ""
            text = para.text.strip()
            if not text:
                continue
            if style.lower().startswith("heading") or style.lower() in {"title", "subtitle"}:
                if style.lower() in {"title", "heading 1", "heading 2"} or style.lower().startswith("heading"):
                    section = text[:280]
                append(f"## {text}")
            else:
                append(text)
        elif child.tag == qn("w:tbl"):
            table = Table(child, doc)
            rows = []
            for row in table.rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                append("\n".join(rows))
    doc_text = "\n\n".join(f"## {s}\n{t}" if s and not t.startswith("## ") else t for s, t in blocks)
    # Consolidar todo el docx como una sola "página" lógica, conservando secciones
    if doc_text.strip():
        result.pages.append(PageContent(page=None, text=doc_text))
    result.page_count = 1
    result.word_count = _count_words(doc_text)
    if cb:
        cb(90, "Documento de Word procesado")
    return result


def _extract_pptx(path: Path, cb: ProgressCb | None) -> ExtractionResult:
    from pptx import Presentation

    result = ExtractionResult()
    prs = Presentation(str(path))
    slides = list(prs.slides)
    result.page_count = len(slides)
    for i, slide in enumerate(slides):
        if cb:
            cb(int((i / max(len(slides), 1)) * 90), f"Leyendo diapositiva {i + 1} de {len(slides)}")
        parts: list[str] = []
        title = None
        try:
            if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
                title = slide.shapes.title.text_frame.text.strip()
        except Exception:  # noqa: BLE001
            title = None
        if title:
            parts.append(f"## {title}")
        for shape in slide.shapes:
            try:
                if shape.has_text_frame:
                    text = shape.text_frame.text.strip()
                    if text and text != title:
                        parts.append(text)
                if getattr(shape, "has_table", False) and shape.has_table:
                    table = shape.table
                    rows = []
                    for row in table.rows:
                        cells = [cell.text.strip() for cell in row.cells]
                        if any(cells):
                            rows.append(" | ".join(cells))
                    if rows:
                        parts.append("\n".join(rows))
            except Exception:  # noqa: BLE001
                continue
        try:
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    parts.append(f"Notas del presentador: {notes}")
        except Exception:  # noqa: BLE001
            pass
        text = "\n".join(parts).strip()
        if text:
            result.pages.append(PageContent(page=i + 1, text=text, section=title))
    result.word_count = sum(_count_words(p.text) for p in result.pages)
    return result


def _extract_txt(path: Path, cb: ProgressCb | None) -> ExtractionResult:
    result = ExtractionResult()
    raw = path.read_bytes()
    for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode("utf-8", errors="replace")
    if cb:
        cb(80, "Archivo de texto leído")
    result.pages.append(PageContent(page=None, text=text.strip()))
    result.page_count = 1
    result.word_count = _count_words(text)
    return result


def _extract_image(path: Path, cb: ProgressCb | None) -> ExtractionResult:
    result = ExtractionResult()
    if cb:
        cb(10, "Aplicando OCR a la imagen")
    mime = "image/png" if path.suffix.lower() == ".png" else "image/jpeg"
    text = _ocr_image_bytes(path.read_bytes(), mime, "imagen de apuntes")
    result.ocr_used = True
    if text:
        result.pages.append(PageContent(page=1, text=text))
    else:
        result.warnings.append("La imagen no contiene texto legible (OCR sin resultados).")
    result.page_count = 1
    result.word_count = _count_words(text)
    if cb:
        cb(80, "OCR completado")
    return result


def extract_document(path: Path, filetype: str, cb: ProgressCb | None = None) -> ExtractionResult:
    """Punto de entrada único de extracción según el tipo de archivo."""
    ft = filetype.lower().lstrip(".")
    if ft == "pdf":
        return _extract_pdf(path, cb)
    if ft == "docx":
        return _extract_docx(path, cb)
    if ft == "pptx":
        return _extract_pptx(path, cb)
    if ft == "txt":
        return _extract_txt(path, cb)
    if ft in {"jpg", "jpeg", "png"}:
        return _extract_image(path, cb)
    raise ValueError(f"Formato no soportado: {ft}")
